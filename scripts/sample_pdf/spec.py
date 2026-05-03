"""Load HQMS field standard from `病案标准.xlsx`.

The spreadsheet is the source of truth for field names, types, lengths,
and required-ness. We cache the parsed result so callers can hit
`load_field_spec()` repeatedly without paying re-parse cost.
"""
from __future__ import annotations

import os
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import openpyxl


@dataclass(frozen=True)
class FieldSpec:
    code: str             # HQMS short code, e.g., 'BAH'
    name: str             # 中文字段含义, e.g., '病案号'
    dtype: str            # '字符' / '数字' / '日期' / '集合'
    length: str | None    # raw length cell as string ('22', '(10,2)', etc.)
    required: bool        # True if 必填; False for None / 条件必填 (treated as optional here)
    note: str | None      # 字段说明 (may reference RC dictionaries)


def _default_spec_path() -> Path:
    here = Path(__file__).resolve()
    # scripts/sample_pdf/ -> repo root
    repo = here.parent.parent.parent
    return repo / "docs" / "病案首页与质控业务" / "病案标准.xlsx"


@lru_cache(maxsize=4)
def load_field_spec(path: str | None = None) -> dict[str, FieldSpec]:
    """Parse the xlsx and return a dict keyed by HQMS code.

    Codes that span multiple slots in the form (e.g. 'QTZD1-QTZD40' or
    'ZZ1-ZZ7') are stored once under the original range string. Callers
    that need to expand them can split on '-' and substitute the index.
    """
    p = Path(path or os.environ.get("SAMPLE_PDF_SPEC") or _default_spec_path())
    if not p.exists():
        raise FileNotFoundError(f"HQMS spec not found at {p}")
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out: dict[str, FieldSpec] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            continue  # header
        idx, code, name, dtype, length, required, note, *_ = list(row) + [None] * 9
        if not code or not name:
            continue
        out[str(code).strip()] = FieldSpec(
            code=str(code).strip(),
            name=str(name).strip(),
            dtype=str(dtype).strip() if dtype else "",
            length=str(length).strip() if length else None,
            required=(str(required).strip() == "必填") if required else False,
            note=str(note).strip() if note else None,
        )
    return out


# Mapping V1 schema fields -> HQMS code, used by mock generators and the
# round-trip integration test ("the field LLM extracts must match what
# the schema expects").
V1_TO_HQMS: dict[str, str] = {
    "recordNo": "BAH",
    "name": "XM",
    "gender": "XB",
    "birthDate": "CSRQ",
    "age": "NL",
    "idCardMasked": "SFZH",
    "admissionDate": "RYSJ",
    "dischargeDate": "CYSJ",
    "lengthOfStay": "SJZY",
    "admissionDept": "RYKB",
    "dischargeDept": "CYKB",
    "admissionRoute": "RYTJ",
    "dischargeStatus": "LYFS",
    "mainDiagnosisCode": "ZYZD_JBBM",
    "mainDiagnosisName": "ZYZD",
    "pathologicalDiagnosis": "BLZD",
    "mainOperationCode": "SSJCZBM1",
    "mainOperationName": "SSJCZMC1",
    "operationDate": "SSJCZRQ1",
    "operator": "SZ1",
    "anesthesiaMethod": "MZFS1",
    "totalCost": "ZFY",
    "drugCost": "XYF",
    "operationCost": "SSF",
    "medicalServiceCost": "YLFWF",
}
